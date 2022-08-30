from docx import Document

class Gather:
    code = {"CS8451":"DAA", "MA8402":"PQT"}
    subjDepartment = {"DAA":"CSE/IT", "PQT":"CSE"}  
    def __init__(self, subject, exam, timetable):
        self.subject = subject
        self.exam = exam
        self.timetable = timetable
        self.tt = Document(self.timetable)
    
    def getDate(self): 
        table = self.tt.tables[0]
        for j in range(2,12):
            for i in range(1,7):
                if table.cell(i,j).text == Gather.code[self.subject]:
                    self.date = table.cell(i,1).text
        return self.date
    
    def getTime(self):
        if int(self.exam[-1]) in [1,2]:
            return "1.30hours"
        else:
            return "3 hours"
    
    def getTotal(self):
        if int(self.exam[-1]) in [1,2]:
            return "50 marks"
        else: 
            return "100 marks"

if __name__ == "__main__":
    obj = Gather("MA8402", "IA2", "TimeTable.docx")
    print(obj.getDate())
    print(obj.getTime())
    print(obj.getTotal())
